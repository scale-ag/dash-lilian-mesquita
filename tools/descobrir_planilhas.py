#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Reconhecimento READ-ONLY das planilhas do cliente (Fase 0) — v4.

Descobertas dos passes anteriores:
  - as duas planilhas SAO publicas (link -> Leitor); o runner le sem credencial;
  - descobrir gid pelo HTML (/htmlview, /edit) NAO funciona nestas planilhas —
    o markup nao traz o menu de abas;
  - o endpoint gviz por NOME de aba funciona.

Entao: os nomes das abas saem do export XLSX (openpyxl), e o conteudo sai do
gviz com headers=0, que devolve TODAS as linhas como dados — sem fundir os
cabecalhos de varias linhas numa so, como o gviz faz por padrao. O gid de cada
aba e' procurado no HTML do /edit por proximidade ao nome (melhor esforco, so
para referencia).

Somente leitura: apenas GETs em endpoints publicos de export. Nunca escreve.
"""
from __future__ import annotations

import io
import json
import pathlib
import re
import urllib.error
import urllib.parse
import urllib.request

from openpyxl import load_workbook

UA = {"User-Agent": "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 "
                    "(KHTML, like Gecko) Chrome/124.0 Safari/537.36"}

PLANILHAS = [
    ("p1_extracao", "1vZgI8ju2OcQit2oEEGPbK-pm19gulnpiFH91TEh3ecI"),
    ("p2_controle", "1ESPchuMZHmXrDIyl5N8Kzy9i20Et0-9EkDVXe_DhSNs"),
]


def get(url: str, timeout: int = 180):
    req = urllib.request.Request(url, headers=UA)
    try:
        with urllib.request.urlopen(req, timeout=timeout) as r:
            return r.status, r.read()
    except urllib.error.HTTPError as e:
        return e.code, e.read()[:2000]
    except Exception as e:                                  # noqa: BLE001
        return 0, f"ERRO {type(e).__name__}: {e}".encode()


def nomes_das_abas(sid: str) -> list[str]:
    st, raw = get(f"https://docs.google.com/spreadsheets/d/{sid}/export?format=xlsx")
    print(f"    export xlsx -> HTTP {st} | {len(raw)} bytes")
    if st != 200 or len(raw) < 1000:
        return []
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        return list(wb.sheetnames)
    except Exception as e:                                  # noqa: BLE001
        print(f"    !! openpyxl falhou: {type(e).__name__}: {e}")
        return []


def gids_por_nome(sid: str, nomes: list[str]) -> dict[str, str]:
    """Melhor esforco: acha o gid citado perto do nome da aba no HTML do /edit."""
    st, raw = get(f"https://docs.google.com/spreadsheets/d/{sid}/edit")
    html = raw.decode("utf-8", "replace")
    print(f"    /edit p/ gids -> HTTP {st} | {len(html)} chars")
    out: dict[str, str] = {}
    for nome in nomes:
        alvo = json.dumps(nome, ensure_ascii=False)[1:-1]      # nome como aparece escapado no JS
        for m in re.finditer(re.escape(alvo), html):
            janela = html[max(0, m.start() - 260): m.end() + 260]
            cand = re.findall(r"\b(\d{6,12})\b", janela) or re.findall(r"\b(0)\b", janela)
            if cand:
                out[nome] = cand[0]
                break
    return out


def csv_da_aba(sid: str, nome: str) -> tuple[int, str]:
    """gviz com headers=0: devolve TODAS as linhas como dados (preserva cabecalho
    de varias linhas, que o gviz fundiria numa so por padrao)."""
    url = (f"https://docs.google.com/spreadsheets/d/{sid}/gviz/tq"
           f"?tqx=out:csv&headers=0&sheet={urllib.parse.quote(nome)}")
    st, raw = get(url)
    if st != 200 or b"accounts.google.com" in raw[:4000]:
        return st, ""
    return st, raw.decode("utf-8", "replace")


def main() -> int:
    pacote: dict = {}
    for chave, sid in PLANILHAS:
        print(f"\n=== {chave} ({sid}) ===")
        nomes = nomes_das_abas(sid)
        print(f"    abas ({len(nomes)}): {nomes}")
        gids = gids_por_nome(sid, nomes) if nomes else {}
        print(f"    gids (melhor esforco): {gids}")
        pacote[chave] = {"sid": sid, "gids": gids, "abas": {}}
        for nome in nomes:
            st, texto = csv_da_aba(sid, nome)
            print(f"    - {nome!r} -> HTTP {st} | {len(texto)} chars")
            pacote[chave]["abas"][nome] = {"http": st, "gid": gids.get(nome), "csv": texto}

    # Grava os CSVs em disco para o job commitar na branch de trabalho: e' assim
    # que eles chegam ao agente, que nao alcanca docs.google.com e precisa deles
    # para montar e testar o build offline.
    destino = pathlib.Path("tools/amostras")
    destino.mkdir(parents=True, exist_ok=True)
    indice = {}
    for chave, sp in pacote.items():
        for nome, info in sp["abas"].items():
            slug = re.sub(r"[^A-Za-z0-9]+", "_", nome).strip("_") or "aba"
            arq = destino / f"{chave}__{slug}.csv"
            arq.write_text(info["csv"], encoding="utf-8")
            indice[f"{chave}/{nome}"] = {"arquivo": arq.name, "gid": info.get("gid"),
                                         "chars": len(info["csv"])}
    (destino / "_indice.json").write_text(
        json.dumps({"planilhas": {k: v["sid"] for k, v in pacote.items()}, "abas": indice},
                   ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"\n{len(indice)} aba(s) gravadas em {destino}/")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
